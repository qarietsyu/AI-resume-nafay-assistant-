"""AI Resume Assistant - ATS score checker built with Streamlit and Gemini."""

import io
import json
import os
import re

import streamlit as st
from docx import Document
from google import genai
from google.genai import types
from pypdf import PdfReader

# Change the default here, or set GEMINI_MODEL in Streamlit secrets / env vars.
DEFAULT_MODEL = "gemini-3.5-flash"
MAX_RESUME_CHARS = 20000
MAX_JD_CHARS = 8000

BREAKDOWN_LABELS = {
    "keyword_match": "Keyword match",
    "formatting": "Formatting & ATS-friendliness",
    "content_quality": "Content quality",
    "structure": "Structure & sections",
}


# ---------------------------------------------------------------- text extraction
def extract_text_from_pdf(data: bytes) -> str:
    reader = PdfReader(io.BytesIO(data))
    return "\n".join((page.extract_text() or "") for page in reader.pages)


def extract_text_from_docx(data: bytes) -> str:
    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def extract_resume_text(filename: str, data: bytes) -> str:
    name = filename.lower()
    if name.endswith(".pdf"):
        text = extract_text_from_pdf(data)
    elif name.endswith(".docx"):
        text = extract_text_from_docx(data)
    elif name.endswith(".txt"):
        text = data.decode("utf-8", errors="ignore")
    else:
        raise ValueError("Unsupported file type. Please upload a PDF, DOCX or TXT file.")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    return text


# ---------------------------------------------------------------- prompt + parsing
def build_prompt(resume_text: str, job_description: str = "") -> str:
    jd_block = (
        f"JOB DESCRIPTION:\n{job_description[:MAX_JD_CHARS]}\n"
        if job_description.strip()
        else "JOB DESCRIPTION: (not provided - evaluate against general ATS best practices)\n"
    )
    return f"""You are an expert ATS (Applicant Tracking System) analyst and resume coach.
Evaluate the resume below{" against the job description" if job_description.strip() else ""}.

Return ONLY a valid JSON object (no markdown, no commentary) with exactly this shape:
{{
  "ats_score": <integer 0-100>,
  "score_breakdown": {{
    "keyword_match": <integer 0-100>,
    "formatting": <integer 0-100>,
    "content_quality": <integer 0-100>,
    "structure": <integer 0-100>
  }},
  "summary": "<2-3 sentence overall assessment>",
  "strengths": ["<short point>", ...],
  "weaknesses": ["<short point>", ...],
  "missing_keywords": ["<keyword or skill>", ...],
  "improvements": [
    {{"section": "<resume section>", "issue": "<what is wrong>", "suggestion": "<specific fix, with example wording>"}}
  ]
}}

Scoring guidance: be realistic and strict; most resumes score between 40 and 85.
Judge keyword relevance, quantified achievements, action verbs, section structure
(contact info, summary, experience, education, skills), clarity, and length.
Give 5-8 concrete improvements.

{jd_block}
RESUME:
{resume_text[:MAX_RESUME_CHARS]}
"""


def _clamp(value, default=0) -> int:
    try:
        return max(0, min(100, int(round(float(value)))))
    except (TypeError, ValueError):
        return default


def _str_list(value) -> list:
    if not isinstance(value, list):
        return []
    return [str(v).strip() for v in value if str(v).strip()]


def parse_response(raw: str) -> dict:
    """Turn the model's reply into a clean, validated result dict."""
    if not raw or not raw.strip():
        raise ValueError("The AI returned an empty response. Please try again.")
    cleaned = re.sub(r"```(?:json)?", "", raw).strip()
    start, end = cleaned.find("{"), cleaned.rfind("}")
    if start == -1 or end <= start:
        raise ValueError("Could not read the AI response. Please try again.")
    try:
        data = json.loads(cleaned[start : end + 1])
    except json.JSONDecodeError as exc:
        raise ValueError("The AI response was not valid JSON. Please try again.") from exc
    if not isinstance(data, dict):
        raise ValueError("Unexpected AI response format. Please try again.")

    breakdown_raw = data.get("score_breakdown")
    breakdown_raw = breakdown_raw if isinstance(breakdown_raw, dict) else {}
    improvements = []
    for item in data.get("improvements") or []:
        if isinstance(item, dict):
            improvements.append(
                {
                    "section": str(item.get("section", "General")).strip() or "General",
                    "issue": str(item.get("issue", "")).strip(),
                    "suggestion": str(item.get("suggestion", "")).strip(),
                }
            )
        elif isinstance(item, str) and item.strip():
            improvements.append({"section": "General", "issue": "", "suggestion": item.strip()})

    return {
        "ats_score": _clamp(data.get("ats_score")),
        "score_breakdown": {k: _clamp(breakdown_raw.get(k)) for k in BREAKDOWN_LABELS},
        "summary": str(data.get("summary", "")).strip(),
        "strengths": _str_list(data.get("strengths")),
        "weaknesses": _str_list(data.get("weaknesses")),
        "missing_keywords": _str_list(data.get("missing_keywords")),
        "improvements": improvements,
    }


# ---------------------------------------------------------------- Gemini
def get_secret(name: str, default: str = "") -> str:
    try:
        if name in st.secrets:
            return str(st.secrets[name])
    except Exception:  # no secrets.toml available
        pass
    return os.environ.get(name, default)


def analyze_resume(api_key: str, model: str, resume_text: str, job_description: str = "") -> dict:
    client = genai.Client(api_key=api_key)
    response = client.models.generate_content(
        model=model,
        contents=build_prompt(resume_text, job_description),
        config=types.GenerateContentConfig(
            temperature=0.2,
            response_mime_type="application/json",
        ),
    )
    return parse_response(response.text)


# ---------------------------------------------------------------- UI
def score_label(score: int) -> str:
    if score >= 80:
        return "Excellent"
    if score >= 65:
        return "Good"
    if score >= 50:
        return "Needs work"
    return "Poor"


def render_results(result: dict) -> None:
    score = result["ats_score"]
    col1, col2 = st.columns([1, 2])
    with col1:
        st.metric("ATS Score", f"{score}/100", score_label(score), delta_color="off")
    with col2:
        st.progress(score / 100)
        if result["summary"]:
            st.write(result["summary"])

    st.subheader("Score breakdown")
    cols = st.columns(len(BREAKDOWN_LABELS))
    for col, (key, label) in zip(cols, BREAKDOWN_LABELS.items()):
        col.metric(label, f"{result['score_breakdown'][key]}/100")

    left, right = st.columns(2)
    with left:
        st.subheader("Strengths")
        for s in result["strengths"] or ["No strengths listed."]:
            st.markdown(f"- {s}")
    with right:
        st.subheader("Weaknesses")
        for w in result["weaknesses"] or ["No weaknesses listed."]:
            st.markdown(f"- {w}")

    if result["missing_keywords"]:
        st.subheader("Missing keywords")
        st.write(", ".join(f"`{k}`" for k in result["missing_keywords"]))

    st.subheader("Suggested improvements")
    for i, imp in enumerate(result["improvements"], 1):
        with st.expander(f"{i}. {imp['section']}", expanded=(i <= 3)):
            if imp["issue"]:
                st.markdown(f"**Issue:** {imp['issue']}")
            if imp["suggestion"]:
                st.markdown(f"**Fix:** {imp['suggestion']}")

    st.download_button(
        "Download report (JSON)",
        data=json.dumps(result, indent=2),
        file_name="ats_report.json",
        mime="application/json",
    )


def main() -> None:
    st.set_page_config(page_title="AI Resume Assistant", page_icon="📄", layout="wide")
    st.title("📄 AI Resume Assistant")
    st.caption("Upload your resume to get an ATS score and tips to improve it.")

    api_key = get_secret("GEMINI_API_KEY")
    model = get_secret("GEMINI_MODEL", DEFAULT_MODEL)

    with st.sidebar:
        st.header("Settings")
        if not api_key:
            api_key = st.text_input("Gemini API key", type="password")
            st.caption("Get a free key at https://aistudio.google.com/apikey")
        model = st.text_input("Gemini model", value=model)

    uploaded = st.file_uploader("Upload your resume", type=["pdf", "docx", "txt"])
    job_description = st.text_area(
        "Job description (optional, improves keyword matching)", height=150
    )

    if st.button("Analyze resume", type="primary", disabled=uploaded is None):
        if not api_key:
            st.error("Please provide a Gemini API key in the sidebar.")
            return
        try:
            with st.spinner("Reading your resume..."):
                text = extract_resume_text(uploaded.name, uploaded.getvalue())
            if len(text) < 100:
                st.error(
                    "Couldn't read enough text from this file. If it's a scanned "
                    "image PDF, upload a text-based PDF or DOCX instead."
                )
                return
            with st.spinner("Analyzing with Gemini..."):
                result = analyze_resume(api_key, model, text, job_description)
        except ValueError as exc:
            st.error(str(exc))
            return
        except Exception as exc:
            st.error(f"Something went wrong: {exc}")
            return
        render_results(result)


if __name__ == "__main__":
    main()
